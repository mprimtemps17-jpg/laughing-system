import json
from pathlib import Path

import pandas as pd
import yaml

from frictionlab.care.analysis import analyze_care
from frictionlab.care.config import CareConfig, load_care_config
from frictionlab.care.humaneval import agreement, export_sheet, import_sheet
from frictionlab.care.items import load_items, render_context
from frictionlab.care.prompts import judge_prompt, organize_prompt
from frictionlab.care.runner import plan_jobs, run_care

ROOT = Path(__file__).resolve().parents[1]
PILOT = ROOT / "data/care/v0.2_pilot.jsonl"


def test_user_materials_load():
    pilot = load_items(PILOT)
    dev = load_items(ROOT / "data/care/v0.1_dev.jsonl")
    assert [it.id for it in pilot] == ["P1-1", "P1-2", "P1-3", "P2-1", "P2-2", "P2-3"]
    assert len(dev) == 20
    # ①と③は対象発言が同一
    assert pilot[0].target == pilot[2].target


def test_stage1_never_sees_target_or_labels():
    it = load_items(PILOT)[0]
    for kind in ("knowledge", "general"):
        _, user, _ = organize_prompt(it, kind)
        assert it.target not in user
        assert "あり" not in user and (it.intended_reason or "x") not in user


def test_judge_prompt_identical_across_conditions_except_table():
    it = load_items(PILOT)[0]
    crit = (ROOT / "configs/care/criteria_v0.2_time.md").read_text(encoding="utf-8")
    _, a, _ = judge_prompt(it, crit)
    _, b, _ = judge_prompt(it, crit, {"constraint": "x"}, "knowledge")
    assert b.startswith(a.split("## 回答方法")[0].rstrip())
    assert it.intended_label not in render_context(it)


def test_alternating_order():
    cfg = CareConfig(name="t", data=str(PILOT), conditions=["A_direct", "B_knowledge", "C_general"])
    jobs = plan_jobs(cfg, load_items(PILOT)[:2])
    assert [j["condition"] for j in jobs] == ["A_direct", "B_knowledge", "C_general",
                                               "B_knowledge", "C_general", "A_direct"]


def test_end_to_end_mock(tmp_path, monkeypatch):
    monkeypatch.chdir(ROOT)
    cfg = load_care_config(ROOT / "configs/care/pilot.yaml")
    out = run_care(cfg, str(ROOT / "configs/care/pilot.yaml"), out_root=str(tmp_path))
    recs = [json.loads(x) for x in (out / "judgments.jsonl").read_text(encoding="utf-8").splitlines()]
    assert len(recs) == 6 * 3
    assert all("stage1" in r for r in recs if r["condition"] != "A_direct")
    report = analyze_care(out, use_intended=True).read_text(encoding="utf-8")
    assert "mockバックエンド" in report and "作成意図" in report


def test_manual_backend_two_stage(tmp_path, monkeypatch):
    monkeypatch.chdir(ROOT)
    raw = yaml.safe_load((ROOT / "configs/care/pilot_manual.yaml").read_text(encoding="utf-8"))
    raw.update(conditions=["B_knowledge"], limit=1)
    p = tmp_path / "c.yaml"
    p.write_text(yaml.safe_dump(raw, allow_unicode=True), encoding="utf-8")
    out = run_care(load_care_config(p), str(p), out_root=str(tmp_path))
    m = out / "manual"
    (m / "P1-1_B_knowledge_0_stage1.response.txt").write_text(
        '```json\n{"constraint":"a","a_knows":"b","b_only":"c","shared":"d","unknown":"e"}\n```', encoding="utf-8")
    run_care(load_care_config(p), str(p), resume=str(out))
    assert (m / "P1-1_B_knowledge_0_final.prompt.md").exists()
    (m / "P1-1_B_knowledge_0_final.response.txt").write_text(
        '{"label":"あり","evidence":[],"reason":"r"}', encoding="utf-8")
    run_care(load_care_config(p), str(p), resume=str(out))
    rec = json.loads((out / "judgments.jsonl").read_text(encoding="utf-8").splitlines()[-1])
    assert rec["label"] == "あり" and rec["stage1"]["parsed"]["constraint"] == "a"


def test_human_eval_roundtrip(tmp_path):
    items = load_items(PILOT)
    for it in items:  # 実データの判定（本人など）は除いて試す
        it.human_labels.clear()
    sheet = export_sheet(items, tmp_path / "s.csv", seed=3)
    df = pd.read_csv(sheet, encoding="utf-8-sig", dtype=str)
    assert "intended" not in " ".join(df.columns) and "scene" not in " ".join(df.columns)
    assert not any(it.intended_reason in " ".join(df["文脈"]) for it in items)
    df["判定"] = ["あり"] * len(df)
    df.to_csv(sheet, index=False, encoding="utf-8-sig")
    items, problems = import_sheet(items, sheet, "x")
    items, _ = import_sheet(items, sheet, "y")
    assert problems == []
    _, table = agreement(items)
    assert table.unanimous.all()


def test_xlsx_roundtrip(tmp_path):
    from openpyxl import load_workbook

    from frictionlab.care.workbook import export_workbook

    items = load_items(PILOT)
    crit = (ROOT / "configs/care/criteria_v0.2_time.md").read_text(encoding="utf-8")
    path = export_workbook(items, tmp_path / "s.xlsx", crit, seed=2)
    wb = load_workbook(path)
    assert wb.sheetnames == ["はじめに", "判定基準", "評価票", "作業記録", "回答データ", "回答の比較"]
    assert wb["回答データ"].sheet_state == "veryHidden"
    flat = " ".join(str(c.value) for row in wb["評価票"].iter_rows() for c in row if c.value)
    assert not any(it.intended_reason in flat for it in items) and "known_" not in flat
    assert not any(it.id in flat for it in items)  # 元のID（変種の番号）を見せない
    for r in range(2, 8):
        wb["評価票"].cell(row=r, column=5, value="なし")
    wb.save(path)
    items, problems = import_sheet(items, path, "z")
    assert problems == [] and all(it.human_labels["z"] == "なし" for it in items)


def test_xlsx_prefill_and_comparison(tmp_path):
    from openpyxl import load_workbook

    from frictionlab.care.workbook import export_workbook

    items = load_items(PILOT)
    for it in items:
        it.human_labels = {"甲": "あり", "乙": "なし"}
        it.human_notes = {"甲": {"reason": "理由甲"}}
    crit = (ROOT / "configs/care/criteria_v0.2_time.md").read_text(encoding="utf-8")
    wb = load_workbook(export_workbook(items, tmp_path / "s.xlsx", crit, seed=1, annotator="甲"))
    assert [wb["評価票"].cell(row=r, column=5).value for r in range(2, 8)] == ["あり"] * 6
    data = [c.value for c in wb["回答データ"][1]]
    assert data == ["ID", "乙の判定", "乙の理由"]  # 自分（甲）は比較相手に含めない
    assert "全件記入後に表示" in wb["回答の比較"]["E7"].value
