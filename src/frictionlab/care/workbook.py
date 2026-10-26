"""評価者に渡す Excel 評価票（.xlsx）を作る。

シート構成
- はじめに : 目的・ルール・記入方法・記入例・進み具合（評価者名もここに書く）
- 判定基準 : 判定基準の全文
- 評価票   : 判定を記入する表（判定はプルダウンで選ぶ）
- 作業記録 : 作業した日・範囲・所要時間・気づいたこと
- 回答の比較 : 自分の判定を全件記入した後にだけ、ほかの回答者の判定・理由が表示される
- 回答データ : 比較用のほかの回答者のデータ（Excelの画面からは開けない veryHidden シート）
作成意図ラベル・場面ID・変種は含めない。
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from .humaneval import presentation_order
from .items import LABELS, CareItem, render_context

FONT = "Meiryo"  # 日本語を読みやすく表示できるフォント（Windows標準。Macでは近いフォントで表示される）
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # 記入するセル（薄い黄色）
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

SHEET_NAME = "評価票"
HEADERS = ["提示順", "ID", "文脈（背景と会話）", "対象発言（A）", "判定", "根拠となる箇所", "理由・迷った点"]


def _f(bold=False, size=10, color="000000"):
    return Font(name=FONT, bold=bold, size=size, color=color)


def export_workbook(items: list[CareItem], out_path: str | Path, criteria: str, seed: int,
                    title: str = "会話の判定 評価票", annotator: str | None = None,
                    include_others: bool = True) -> Path:
    """annotator を指定すると、その人の記録済みの判定を記入した状態で作り、比較シートからはその人を除く。"""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    ordered = presentation_order(items, seed)
    n = len(ordered)
    wb = Workbook()

    # ---------------------------------------------------------------- はじめに
    ws = wb.active
    ws.title = "はじめに"
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 90
    rows = [
        (title, None, "title"),
        ("", None, None),
        ("評価者名", annotator or "", "input"),
        ("記入済み", f'=COUNTA({SHEET_NAME}!E2:E{n + 1})&" / {n} 件"', "formula"),
        ("", None, None),
        ("この評価票について", "短い日本語の会話を読み、Aの「対象発言」について判定してもらいます。研究で使う正解を決めるための、大切な判定です。", None),
        ("ルール", "・ほかの人と相談せず、一人で判定してください（他の人の回答やAIの回答も見ないでください）。\n"
                   "・「判定基準」シートを先に読んでから始めてください。\n"
                   "・本文に書かれていることだけで判断してください。登場人物の性格や気持ちは想像しなくて大丈夫です。\n"
                   "・1件ずつ独立に判断してください。似た会話があっても、前の答えに合わせる必要はありません。\n"
                   "・迷ったときは、迷った理由を書いてもらえるととても助かります。", None),
        ("記入するところ", "「評価票」シートの黄色いセル（判定・根拠となる箇所・理由・迷った点）と、このシートの評価者名です。\n"
                         "判定はセルをクリックすると出る ▼ から選びます（あり／なし／判定対象外／保留）。", None),
        ("作業記録", "作業した日ごとに「作業記録」シートに1行書いてください（所要時間は研究の計画に使います）。", None),
        ("回答の比較", "全件の判定を記入し終えると、「回答の比較」シートに、ほかの回答者の判定と理由が表示されます。"
                      "自分の判定をすべて終えてから開いてください（先に見ても、全件記入するまでは表示されません）。", None),
        ("終わったら", "このファイルを保存して、研究者に返してください。", None),
        ("", None, None),
        ("記入例", "（この例は説明用です。評価票の会話とは関係ありません）", "bold"),
    ]
    for i, (k, v, kind) in enumerate(rows, 1):
        a = ws.cell(row=i, column=1, value=k)
        a.font = _f(bold=True)
        if kind == "title":
            a.font = _f(bold=True, size=14)  # 折り返さず右の空セルにはみ出させる
            continue
        a.alignment = WRAP
        if v is None:
            continue
        b = ws.cell(row=i, column=2, value=v)
        b.font = _f(bold=kind == "bold")
        b.alignment = WRAP
        if kind == "input":
            b.fill = INPUT_FILL
            b.border = BOX
        if v and kind is None:
            ws.row_dimensions[i].height = 15 * (v.count("\n") + 1 + len(v) // 60)
    ex = [
        ("文脈", "【背景】\nAとBは同じゼミの学生。一緒に昼ごはんを食べる時刻を決めている。\n\n【会話】\nB「明日は10時から12時までゼミがある」\nA「了解」\nB「昼ごはんの時刻を決めよう」"),
        ("対象発言（A）", "ゼミが終わった後の12時半からならどう？"),
        ("判定", "なし"),
        ("根拠となる箇所", "B「明日は10時から12時までゼミがある」A「了解」"),
        ("理由・迷った点", "Aはゼミの時間を聞いていて、終わった後の時刻で確認している。"),
    ]
    for j, (k, v) in enumerate(ex, len(rows) + 1):
        ws.cell(row=j, column=1, value=k).font = _f(bold=True)
        c = ws.cell(row=j, column=2, value=v)
        c.font = _f(color="595959")
        c.alignment = WRAP
        c.border = BOX
        ws.row_dimensions[j].height = 15 * (v.count("\n") + 1)

    # ---------------------------------------------------------------- 判定基準
    wc = wb.create_sheet("判定基準")
    wc.column_dimensions["A"].width = 110
    wc.cell(row=1, column=1, value="判定基準").font = _f(bold=True, size=13)
    for i, line in enumerate(criteria.strip().splitlines(), 3):
        c = wc.cell(row=i, column=1, value=line)
        c.font = _f()
        c.alignment = Alignment(wrap_text=True)

    # ---------------------------------------------------------------- 評価票
    we = wb.create_sheet(SHEET_NAME)
    widths = [7, 9, 60, 34, 13, 34, 40]
    for col, (h, w) in enumerate(zip(HEADERS, widths), 1):
        c = we.cell(row=1, column=col, value=h)
        c.font = _f(bold=True)
        c.fill = HEAD_FILL
        c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        we.column_dimensions[c.column_letter].width = w
    we.freeze_panes = "C2"
    dv = DataValidation(type="list", formula1='"' + ",".join(LABELS) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="判定", error="あり／なし／判定対象外／保留 から選んでください")
    we.add_data_validation(dv)
    for r, it in enumerate(ordered, 2):
        ctx = render_context(it)
        mine = annotator and it.human_labels.get(annotator)
        note = it.human_notes.get(annotator, {}) if mine else {}
        vals = [r - 1, it.id, ctx, it.target, mine or None, note.get("evidence") or None, note.get("reason") or None]
        for col, v in enumerate(vals, 1):
            c = we.cell(row=r, column=col, value=v)
            c.font = _f()
            c.alignment = WRAP
            c.border = BOX
            if col >= 5:
                c.fill = INPUT_FILL
        we.cell(row=r, column=4).font = _f(bold=True)
        dv.add(we.cell(row=r, column=5))
        we.row_dimensions[r].height = max(90, 15 * (ctx.count("\n") + 2))

    # 印刷したときに横1ページに収める
    for sheet in (ws, we):
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    we.print_title_rows = "1:1"

    # ---------------------------------------------------------------- 作業記録
    wl = wb.create_sheet("作業記録")
    for col, (h, w) in enumerate(zip(["日付", "判定した範囲（提示順）", "所要時間（分）", "気づいたこと・感想"],
                                     [14, 24, 14, 70]), 1):
        c = wl.cell(row=1, column=col, value=h)
        c.font = _f(bold=True)
        c.fill = HEAD_FILL
        c.border = BOX
        wl.column_dimensions[c.column_letter].width = w
    for r in range(2, 12):
        for col in range(1, 5):
            c = wl.cell(row=r, column=col)
            c.fill = INPUT_FILL
            c.border = BOX
            c.font = _f()
            c.alignment = WRAP

    if include_others:
        _comparison_sheets(wb, ordered, annotator)

    wb.active = wb.sheetnames.index("はじめに")
    # 数式の結果は保存せず、Excelで開いたときに計算させる（LibreOfficeで再計算すると veryHidden が解除されるため）
    wb.calculation.fullCalcOnLoad = True
    wb.properties.title = title
    wb.save(out_path)
    out_path.with_suffix(".order.txt").write_text(
        f"作成日={date.today()} seed={seed}\n" + " → ".join(it.id for it in ordered) + "\n", encoding="utf-8")
    return out_path


def _comparison_sheets(wb: Workbook, ordered: list[CareItem], annotator: str | None) -> None:
    """自分の判定を全件記入した後にだけ、ほかの回答者の判定を表示する比較シートを作る。"""
    n = len(ordered)
    others = sorted({a for it in ordered for a in it.human_labels if a != annotator})
    q = f"'{SHEET_NAME}'"

    # 比較用データ（Excelの画面からは開けない）。列: ID, 回答者1の判定, 回答者1の理由, 回答者2の判定, ...
    wd = wb.create_sheet("回答データ")
    wd.append(["ID"] + [x for o in others for x in (f"{o}の判定", f"{o}の理由")])
    for it in ordered:
        row = [it.id]
        for o in others:
            note = it.human_notes.get(o, {})
            reason = "／".join(x for x in (note.get("evidence", ""), note.get("reason", "")) if x)
            row += [it.human_labels.get(o, ""), reason]
        wd.append(row)
    wd.sheet_state = "veryHidden"

    wc = wb.create_sheet("回答の比較")
    wc.column_dimensions["A"].width = 7
    wc.column_dimensions["B"].width = 9
    wc.column_dimensions["C"].width = 34
    wc.column_dimensions["D"].width = 13
    wc.cell(row=1, column=1, value="回答の比較").font = _f(bold=True, size=14)
    wc.cell(row=2, column=1, value="状態").font = _f(bold=True)
    done = "$B$3"
    wc.cell(row=3, column=1, value="全件記入").font = _f(bold=True)
    wc.cell(row=3, column=2, value=f"=COUNTA({q}!$E$2:$E${n + 1})={n}").font = _f()
    msg = wc.cell(row=2, column=2, value=(
        f'=IF({done},"全件の記入が終わりました。ほかの回答者の判定と比べてみてください。",'
        f'"あと "&({n}-COUNTA({q}!$E$2:$E${n + 1}))&" 件記入すると、ほかの回答者の判定が表示されます。")'))
    msg.font = _f(bold=True, color="C00000")
    if not others:
        wc.cell(row=4, column=1, value="（まだほかの回答者の記録はありません）").font = _f(color="595959")

    head = 6
    headers = ["提示順", "ID", "対象発言（A）", "あなたの判定"]
    for o in others:
        headers += [f"{o}の判定", "一致", f"{o}の理由"]
    for col, h in enumerate(headers, 1):
        c = wc.cell(row=head, column=col, value=h)
        c.font = _f(bold=True)
        c.fill = HEAD_FILL
        c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    lock = "（全件記入後に表示）"
    for i in range(n):
        r, src = head + 1 + i, i + 2
        cells = [f"={q}!A{src}", f"={q}!B{src}", f"={q}!D{src}", f'=IF({q}!E{src}="","",{q}!E{src})']
        for k in range(len(others)):
            lab_col, rea_col = 2 + 2 * k, 3 + 2 * k
            look = lambda col: (f"INDEX('回答データ'!$A$2:${_col(col)}${n + 1},"
                                f"MATCH($B{r},'回答データ'!$A$2:$A${n + 1},0),{col})")
            cells += [f'=IF({done},{look(lab_col)},"{lock}")',
                      f'=IF({done},IF($D{r}={_col(5 + 3 * k)}{r},"○","×"),"")',
                      f'=IF({done},{look(rea_col)},"")']
        for col, v in enumerate(cells, 1):
            c = wc.cell(row=r, column=col, value=v)
            c.font = _f(bold=col == 4)
            c.alignment = WRAP
            c.border = BOX
        wc.row_dimensions[r].height = 45
    for k, o in enumerate(others):
        lab = _col(5 + 3 * k)
        wc.column_dimensions[lab].width = 13
        wc.column_dimensions[_col(6 + 3 * k)].width = 6
        wc.column_dimensions[_col(7 + 3 * k)].width = 40
        rr = head + n + 2 + k
        wc.cell(row=rr, column=3, value=f"{o}との一致率").font = _f(bold=True)
        c = wc.cell(row=rr, column=4, value=(
            f'=IF({done},SUMPRODUCT(--($D${head + 1}:$D${head + n}={lab}${head + 1}:{lab}${head + n}))/{n},"")'))
        c.number_format = "0%"
        c.font = _f(bold=True)
    wc.freeze_panes = f"A{head + 1}"


def _col(i: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(i)
